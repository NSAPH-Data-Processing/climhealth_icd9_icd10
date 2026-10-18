"""Step 2.6a: turn each fetched full text and supplement into plain text with location markers.

Input: 024_fulltext_manifest.json (papers with status fulltext / fulltext_no_supplements).
For every file of a paper, writes papers/<slug>/text/<file name>.txt (git-ignored, like papers/):
  PDF   "=== page N ===" before each page. Uses poppler's pdftotext when it is on the PATH (it keeps
        two-column layouts in reading order); otherwise pdfplumber (columns get interleaved)
  XML   "=== section: A > B ===" for JATS sections, "=== table: Table 1. caption ===" for tables
        (rows tab-separated, footnotes kept), "=== abstract ===", "=== back: ... ===" for
        appendices / footnotes / supplementary-material captions
  DOCX  paragraphs in order, "=== table N ===" for each table (cells tab-separated)
  XLSX  "=== sheet: name ===", one row per line
Zip supplements are not extracted: the two in scope hold R code and figure source data only.
These markers are what Step 2.6 evidence locations refer to.

Run: .venv/bin/python 026_extract_text.py
"""
import json
import os
import re
import shutil
import subprocess
from xml.etree import ElementTree

import docx
import openpyxl
import pdfplumber

IN = "024_fulltext_manifest.json"


def norm(s):
    return re.sub(r"\s+", " ", s or "").strip()


def itertext(el):
    return norm("".join(el.itertext()))


def jats_text(path):
    root = ElementTree.parse(path).getroot()
    out = []

    def table(tw):
        label = itertext(tw.find("label")) if tw.find("label") is not None else ""
        cap = itertext(tw.find("caption")) if tw.find("caption") is not None else ""
        out.append("=== table: %s %s ===" % (label, cap))
        for tr in tw.iter("tr"):
            out.append("\t".join(itertext(c) for c in tr if c.tag in ("td", "th")))
        for foot in tw.iter("table-wrap-foot"):
            out.append(itertext(foot))

    def walk(el, path):
        for ch in el:
            if ch.tag in ("sec", "app"):
                t = ch.find("title")
                title = itertext(t) if t is not None else (itertext(ch.find("label"))
                                                          if ch.find("label") is not None else "")
                p = path + [title] if title else path
                out.append("=== %s: %s ===" % ("section" if ch.tag == "sec" else "back",
                                               " > ".join(p)))
                walk(ch, p)
            elif ch.tag == "table-wrap":
                table(ch)
                out.append("=== section: %s ===" % " > ".join(path))
            elif ch.tag in ("p", "list"):
                # tables nested in paragraphs are written separately, after the paragraph
                tws = list(ch.iter("table-wrap"))
                for parent in list(ch.iter()):
                    for tw in [c for c in parent if c.tag == "table-wrap"]:
                        parent.remove(tw)
                out.append(itertext(ch))
                for tw in tws:
                    table(tw)
            elif ch.tag == "fig":
                out.append("[figure] " + itertext(ch.find("caption") if ch.find("caption") is not None else ch))
            elif ch.tag == "supplementary-material":
                out.append("[supplementary-material] " + itertext(ch))
            elif ch.tag not in ("title", "label", "ref-list"):
                walk(ch, path)

    for ab in root.iter("abstract"):
        out.append("=== abstract ===")
        out.append(itertext(ab))
    body = root.find("body")
    if body is not None:
        walk(body, [])
    back = root.find("back")
    if back is not None:
        for ch in back:
            if ch.tag == "ref-list":
                continue
            if ch.tag == "app-group":
                walk(ch, ["Appendix"])
            else:
                out.append("=== back: %s ===" % ch.tag)
                walk(ch, [ch.tag]) if len(ch) else out.append(itertext(ch))
    floats = root.find("floats-group")  # author manuscripts keep tables and figures here
    if floats is not None:
        walk(floats, ["floats"])
    return "\n".join(l for l in out if l.strip())


def pdf_text(path):
    out = []
    if shutil.which("pdftotext"):
        raw = subprocess.run(["pdftotext", "-enc", "UTF-8", path, "-"], stdout=subprocess.PIPE, stderr=subprocess.DEVNULL,
                             check=True).stdout.decode("utf-8", "replace")
        for i, page in enumerate(raw.split("\f"), 1):
            if page.strip() or i == 1:
                out.append("=== page %d ===" % i)
                out.append(page)
        return "\n".join(out)
    with pdfplumber.open(path) as pdf:
        for i, page in enumerate(pdf.pages, 1):
            out.append("=== page %d ===" % i)
            out.append(page.extract_text() or "")
    return "\n".join(out)


def docx_text(path):
    d = docx.Document(path)
    out, ti = [], 0
    body = d.element.body
    tables = iter(d.tables)
    paras = {p._p: p for p in d.paragraphs}
    for child in body.iterchildren():
        tag = child.tag.split("}")[1]
        if tag == "p" and child in paras:
            out.append(paras[child].text)
        elif tag == "tbl":
            ti += 1
            t = next(tables)
            out.append("=== table %d ===" % ti)
            for row in t.rows:
                cells = []
                for c in row.cells:
                    txt = norm(c.text)
                    if not cells or cells[-1] != txt:  # merged cells repeat
                        cells.append(txt)
                out.append("\t".join(cells))
    return "\n".join(l for l in out if l.strip())


def xlsx_text(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        out.append("=== sheet: %s ===" % ws.title)
        for row in ws.iter_rows(values_only=True):
            vals = ["" if v is None else str(v) for v in row]
            if any(vals):
                out.append("\t".join(vals).rstrip())
    return "\n".join(out)


PARSERS = {".pdf": pdf_text, ".xml": jats_text, ".docx": docx_text, ".xlsx": xlsx_text}


def main():
    papers = json.load(open(IN, encoding="utf-8"))["papers"]
    n = 0
    for p in papers:
        files = [f for f in [p["fulltext_xml"], p["fulltext_pdf"]] + p["supplements"] if f]
        tdir = os.path.join("papers", p["slug"], "text")
        os.makedirs(tdir, exist_ok=True) if files else None
        for f in files:
            ext = os.path.splitext(f)[1].lower()
            out = os.path.join(tdir, os.path.basename(f) + ".txt")
            if ext not in PARSERS:
                print("skip (no parser): %s" % f)
                continue
            if os.path.exists(out):
                continue
            try:
                txt = PARSERS[ext](f)
            except Exception as e:  # report and go on
                print("FAILED %s: %r" % (f, e))
                continue
            with open(out, "w", encoding="utf-8") as fh:
                fh.write(txt + "\n")
            n += 1
            print("%7d chars  %s" % (len(txt), out), flush=True)
    print("wrote %d text files" % n)


if __name__ == "__main__":
    main()
